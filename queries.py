from db_connection import DB
from docx import Document


class Queries:
    @staticmethod
    def num_of_cars_brand():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select Marka, count(*) as Broj
            from automobili
            group by Marka""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Spisak svih marki i njihov ukupan broj automobila:\n\n').bold = True

        for res in result:
            p.add_run(res[0] + ' - ' + str(res[1]) + '\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def num_of_cars_location():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select Lokacija, count(*) as Broj
            from automobili
            group by Lokacija""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Spisak svih lokacija i ukupnog broja automobila po lokaciji:\n\n').bold = True

        for res in result:
            p.add_run(res[0] + ' - ' + str(res[1]) + '\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def num_of_cars_color():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select Boja, count(*) as Broj
            from automobili
            group by Boja""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Spisak svih boja i ukupnog broja automobila po boji:\n\n').bold = True

        for res in result:
            p.add_run(res[0] + ' - ' + str(res[1]) + '\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def highest_price():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select *
            from automobili
            order by Cena desc
            limit 30""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Top 30 najskupljih automobila:\n\n').bold = True

        for res in result:
            p.add_run('Stanje: ' + ('/' if res[1] is None else res[1]) + '\n')
            p.add_run('Marka: ' + ('/' if res[2] is None else res[2]) + '\n')
            p.add_run('Model: ' + ('/' if res[3] is None else res[3]) + '\n')
            p.add_run('Godište: ' + ('/' if res[4] is None else (str(res[4]) + '.')) + '\n')
            p.add_run('Kilometraža: ' + ('/' if res[5] is None else (str(res[5]) + ' km')) + '\n')
            p.add_run('Karoserija: ' + ('/' if res[6] is None else res[6]) + '\n')
            p.add_run('Gorivo: ' + ('/' if res[7] is None else res[7]) + '\n')
            p.add_run('Kubikaža: ' + ('/' if res[8] is None else (str(res[8]) + ' cm³')) + '\n')
            p.add_run('Snaga: ' + ('/' if res[9] is None else (str(res[9]) + ' KS')) + '\n')
            p.add_run('Menjač: ' + ('/' if res[10] is None else res[10]) + '\n')
            p.add_run('Vrata: ' + ('/' if res[11] is None else res[11]) + '\n')
            p.add_run('Boja: ' + ('/' if res[12] is None else res[12]) + '\n')
            p.add_run('Lokacija: ' + ('/' if res[13] is None else res[13]) + '\n')
            p.add_run('Cena: ' + ('Po dogovoru' if res[14] == 0 or res[14] is None else (str(res[14]) + ' €')) + '\n')
            p.add_run('\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def highest_price_suv():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select *
            from automobili
            where Karoserija = 'Džip/SUV'
            order by Cena desc
            limit 30""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Top 30 najskupljih džip/SUV automobila:\n\n').bold = True

        for res in result:
            p.add_run('Stanje: ' + ('/' if res[1] is None else res[1]) + '\n')
            p.add_run('Marka: ' + ('/' if res[2] is None else res[2]) + '\n')
            p.add_run('Model: ' + ('/' if res[3] is None else res[3]) + '\n')
            p.add_run('Godište: ' + ('/' if res[4] is None else (str(res[4]) + '.')) + '\n')
            p.add_run('Kilometraža: ' + ('/' if res[5] is None else (str(res[5]) + ' km')) + '\n')
            p.add_run('Gorivo: ' + ('/' if res[7] is None else res[7]) + '\n')
            p.add_run('Kubikaža: ' + ('/' if res[8] is None else (str(res[8]) + ' cm³')) + '\n')
            p.add_run('Snaga: ' + ('/' if res[9] is None else (str(res[9]) + ' KS')) + '\n')
            p.add_run('Menjač: ' + ('/' if res[10] is None else res[10]) + '\n')
            p.add_run('Vrata: ' + ('/' if res[11] is None else res[11]) + '\n')
            p.add_run('Boja: ' + ('/' if res[12] is None else res[12]) + '\n')
            p.add_run('Lokacija: ' + ('/' if res[13] is None else res[13]) + '\n')
            p.add_run('Cena: ' + ('Po dogovoru' if res[14] == 0 or res[14] is None else (str(res[14]) + ' €')) + '\n')
            p.add_run('\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def year_2021_2022():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select *
            from automobili
            where Godiste in (2021, 2022)
            order by Cena desc""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Rang lista automobila proizvedenih 2021/2022:\n\n').bold = True

        for res in result:
            p.add_run('Stanje: ' + ('/' if res[1] is None else res[1]) + '\n')
            p.add_run('Marka: ' + ('/' if res[2] is None else res[2]) + '\n')
            p.add_run('Model: ' + ('/' if res[3] is None else res[3]) + '\n')
            p.add_run('Godište: ' + ('/' if res[4] is None else (str(res[4]) + '.')) + '\n')
            p.add_run('Kilometraža: ' + ('/' if res[5] is None else (str(res[5]) + ' km')) + '\n')
            p.add_run('Karoserija: ' + ('/' if res[6] is None else res[6]) + '\n')
            p.add_run('Gorivo: ' + ('/' if res[7] is None else res[7]) + '\n')
            p.add_run('Kubikaža: ' + ('/' if res[8] is None else (str(res[8]) + ' cm³')) + '\n')
            p.add_run('Snaga: ' + ('/' if res[9] is None else (str(res[9]) + ' KS')) + '\n')
            p.add_run('Menjač: ' + ('/' if res[10] is None else res[10]) + '\n')
            p.add_run('Vrata: ' + ('/' if res[11] is None else res[11]) + '\n')
            p.add_run('Boja: ' + ('/' if res[12] is None else res[12]) + '\n')
            p.add_run('Lokacija: ' + ('/' if res[13] is None else res[13]) + '\n')
            p.add_run('Cena: ' + ('Po dogovoru' if res[14] == 0 or res[14] is None else (str(res[14]) + ' €')) + '\n')
            p.add_run('\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def highest_capacity():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select *
            from automobili
            where Kubikaza in (select MAX(Kubikaza) from automobili)""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Automobil(i) sa najvećim kapacitetom:\n\n').bold = True

        for res in result:
            p.add_run('Stanje: ' + ('/' if res[1] is None else res[1]) + '\n')
            p.add_run('Marka: ' + ('/' if res[2] is None else res[2]) + '\n')
            p.add_run('Model: ' + ('/' if res[3] is None else res[3]) + '\n')
            p.add_run('Godište: ' + ('/' if res[4] is None else (str(res[4]) + '.')) + '\n')
            p.add_run('Kilometraža: ' + ('/' if res[5] is None else (str(res[5]) + ' km')) + '\n')
            p.add_run('Karoserija: ' + ('/' if res[6] is None else res[6]) + '\n')
            p.add_run('Gorivo: ' + ('/' if res[7] is None else res[7]) + '\n')
            p.add_run('Kubikaža: ' + ('/' if res[8] is None else (str(res[8]) + ' cm³')) + '\n')
            p.add_run('Snaga: ' + ('/' if res[9] is None else (str(res[9]) + ' KS')) + '\n')
            p.add_run('Menjač: ' + ('/' if res[10] is None else res[10]) + '\n')
            p.add_run('Vrata: ' + ('/' if res[11] is None else res[11]) + '\n')
            p.add_run('Boja: ' + ('/' if res[12] is None else res[12]) + '\n')
            p.add_run('Lokacija: ' + ('/' if res[13] is None else res[13]) + '\n')
            p.add_run('Cena: ' + ('Po dogovoru' if res[14] == 0 or res[14] is None else (str(res[14]) + ' €')) + '\n')
            p.add_run('\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def highest_engine_power():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select *
            from automobili
            where Snaga in (select MAX(Snaga) from automobili)""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Automobil(i) sa najvećom snagom motora:\n\n').bold = True

        for res in result:
            p.add_run('Stanje: ' + ('/' if res[1] is None else res[1]) + '\n')
            p.add_run('Marka: ' + ('/' if res[2] is None else res[2]) + '\n')
            p.add_run('Model: ' + ('/' if res[3] is None else res[3]) + '\n')
            p.add_run('Godište: ' + ('/' if res[4] is None else (str(res[4]) + '.')) + '\n')
            p.add_run('Kilometraža: ' + ('/' if res[5] is None else (str(res[5]) + ' km')) + '\n')
            p.add_run('Karoserija: ' + ('/' if res[6] is None else res[6]) + '\n')
            p.add_run('Gorivo: ' + ('/' if res[7] is None else res[7]) + '\n')
            p.add_run('Kubikaža: ' + ('/' if res[8] is None else (str(res[8]) + ' cm³')) + '\n')
            p.add_run('Snaga: ' + ('/' if res[9] is None else (str(res[9]) + ' KS')) + '\n')
            p.add_run('Menjač: ' + ('/' if res[10] is None else res[10]) + '\n')
            p.add_run('Vrata: ' + ('/' if res[11] is None else res[11]) + '\n')
            p.add_run('Boja: ' + ('/' if res[12] is None else res[12]) + '\n')
            p.add_run('Lokacija: ' + ('/' if res[13] is None else res[13]) + '\n')
            p.add_run('Cena: ' + ('Po dogovoru' if res[14] == 0 or res[14] is None else (str(res[14]) + ' €')) + '\n')
            p.add_run('\n')

        document.save('word_documents/results.docx')

    @staticmethod
    def greatest_distance():
        conn = DB.get_connection()

        cursor = conn.cursor()

        cursor.execute("""select *
            from automobili
            where Kilometraza in (select MAX(Kilometraza) from automobili)""")

        result = cursor.fetchall()

        document = Document('word_documents/results.docx')

        p = document.add_paragraph()
        p.add_run('Automobil(i) sa najvećom kilometražom:\n\n').bold = True

        for res in result:
            p.add_run('Stanje: ' + ('/' if res[1] is None else res[1]) + '\n')
            p.add_run('Marka: ' + ('/' if res[2] is None else res[2]) + '\n')
            p.add_run('Model: ' + ('/' if res[3] is None else res[3]) + '\n')
            p.add_run('Godište: ' + ('/' if res[4] is None else (str(res[4]) + '.')) + '\n')
            p.add_run('Kilometraža: ' + ('/' if res[5] is None else (str(res[5]) + ' km')) + '\n')
            p.add_run('Karoserija: ' + ('/' if res[6] is None else res[6]) + '\n')
            p.add_run('Gorivo: ' + ('/' if res[7] is None else res[7]) + '\n')
            p.add_run('Kubikaža: ' + ('/' if res[8] is None else (str(res[8]) + ' cm³')) + '\n')
            p.add_run('Snaga: ' + ('/' if res[9] is None else (str(res[9]) + ' KS')) + '\n')
            p.add_run('Menjač: ' + ('/' if res[10] is None else res[10]) + '\n')
            p.add_run('Vrata: ' + ('/' if res[11] is None else res[11]) + '\n')
            p.add_run('Boja: ' + ('/' if res[12] is None else res[12]) + '\n')
            p.add_run('Lokacija: ' + ('/' if res[13] is None else res[13]) + '\n')
            p.add_run('Cena: ' + ('Po dogovoru' if res[14] == 0 or res[14] is None else (str(res[14]) + ' €')) + '\n')
            p.add_run('\n')

        document.save('word_documents/results.docx')
